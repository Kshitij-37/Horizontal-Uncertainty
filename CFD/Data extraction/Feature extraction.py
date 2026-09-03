"""
Wind Data Extractor
-------------------
Scans a parent folder for Stab_X subfolders, reads all .txt files in each,
produces one Excel file per stability class, and one combined Excel file
with all stab sheets + a weighted Net sheet.

Just run:  python extract_wind_data.py
"""

import os
import re
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


ALL_ANGLES = [0, 30, 60, 90, 120, 150, 180, 210, 240, 270, 300, 330]
COLUMNS = ["Vh", "TI", "D (°)", "I (°)"]
DISPLAY_HEADERS = ["Angle (°)", "Vh", "TI", "D (°)", "I (°)"]


# ── File parsing ─────────────────────────────────────────────────────────────

def parse_angle_from_filename(filename):
    basename = os.path.splitext(os.path.basename(filename))[0]
    match = re.search(r'_(\d+)$', basename)
    if not match:
        return None
    angle = int(match.group(1)) // 10
    return 0 if angle == 360 else angle


def parse_txt_file(filepath):
    with open(filepath, 'r', encoding='utf-8-sig') as f:
        content = f.read().replace('\xa0', ' ')
    lines = content.splitlines(keepends=True)

    header_line = None
    data_lines = []
    for line in lines:
        stripped = line.strip()
        if not stripped:
            continue
        if stripped.startswith("Label"):
            header_line = stripped
        elif header_line is not None:
            data_lines.append(stripped)

    if header_line is None or not data_lines:
        return None

    headers = re.split(r'\s{2,}', header_line)
    parsed_rows = []
    for line in data_lines:
        parts = re.split(r'\s{2,}', line)
        if len(parts) < len(headers):
            parts += [None] * (len(headers) - len(parts))
        row = dict(zip(headers, parts))
        if re.match(r'(?i)lidar', row.get("Label", "")):
            continue
        parsed_rows.append(row)

    if not parsed_rows:
        return None

    df = pd.DataFrame(parsed_rows)
    for col in ["H (m)", "Vh", "TI", "D (°)", "I (°)"]:
        if col in df.columns:
            df[col] = pd.to_numeric(df[col], errors='coerce')
    return df


def extract_location_name(label):
    return label.split('#')[0].strip()


def collect_data(folder, target_h):
    """Returns { location: { angle: { col: value } } }"""
    data = {}
    txt_files = [f for f in os.listdir(folder) if f.lower().endswith('.txt')]
    if not txt_files:
        print(f"    No .txt files found in: {folder}")
        return data

    for filename in txt_files:
        angle = parse_angle_from_filename(filename)
        if angle is None:
            print(f"    [SKIP] Could not parse angle from: {filename}")
            continue

        df = parse_txt_file(os.path.join(folder, filename))
        if df is None or df.empty:
            print(f"    [SKIP] No valid data in: {filename}")
            continue

        subset = df[df["H (m)"] == float(target_h)]
        if subset.empty:
            print(f"    [WARN] H={target_h} not found in: {filename}")
            continue

        for _, row in subset.iterrows():
            loc = extract_location_name(str(row.get("Label", "")))
            if not loc:
                continue
            data.setdefault(loc, {})[angle] = {
                "Vh":    row.get("Vh"),
                "TI":    row.get("TI"),
                "D (°)": row.get("D (°)"),
                "I (°)": row.get("I (°)"),
            }
    return data


# ── Weighted net calculation ──────────────────────────────────────────────────

def compute_net_data(all_stab_data, weights):
    """
    all_stab_data: { stab_name: { loc: { angle: { col: val } } } }
    weights:       { stab_name: float }
    Returns:       { loc: { angle: { col: weighted_avg } } }
    """
    total_weight = sum(weights.values())
    net = {}

    # Collect all locations across all stabs
    all_locs = set()
    for stab_data in all_stab_data.values():
        all_locs.update(stab_data.keys())

    for loc in all_locs:
        net[loc] = {}
        for angle in ALL_ANGLES:
            angle_net = {}
            for col in COLUMNS:
                weighted_sum = 0.0
                effective_weight = 0.0
                for stab_name, w in weights.items():
                    val = all_stab_data.get(stab_name, {}).get(loc, {}).get(angle, {}).get(col)
                    if val is not None and not (isinstance(val, float) and pd.isna(val)):
                        weighted_sum += val * w
                        effective_weight += w
                angle_net[col] = (weighted_sum / effective_weight) if effective_weight > 0 else None
            net[loc][angle] = angle_net
    return net


# ── Excel writing ─────────────────────────────────────────────────────────────

def make_styles():
    return {
        "header_font":  Font(name="Arial", bold=True, color="FFFFFF", size=10),
        "header_fill":  PatternFill("solid", start_color="2F5496"),
        "net_fill":     PatternFill("solid", start_color="1D6A3A"),   # green header for Net
        "loc_font":     Font(name="Arial", bold=True, size=10),
        "data_font":    Font(name="Arial", size=10),
        "center":       Alignment(horizontal="center", vertical="center"),
        "thin_border":  Border(
            left=Side(style="thin", color="CCCCCC"),
            right=Side(style="thin", color="CCCCCC"),
            top=Side(style="thin", color="CCCCCC"),
            bottom=Side(style="thin", color="CCCCCC"),
        ),
        "alt_fill":     PatternFill("solid", start_color="EEF2F8"),
        "net_alt_fill": PatternFill("solid", start_color="E6F2EC"),   # light green alt rows
    }


def write_data_sheet(wb, sheet_title, title_text, angle_data, styles, is_net=False):
    sheet = wb.create_sheet(title=sheet_title[:31])

    sheet.merge_cells("A1:E1")
    tc = sheet["A1"]
    tc.value = title_text
    tc.font = Font(name="Arial", bold=True, size=12, color="1D6A3A" if is_net else "2F5496")
    tc.alignment = styles["center"]
    sheet.row_dimensions[1].height = 22

    hdr_fill = styles["net_fill"] if is_net else styles["header_fill"]
    for col_idx, header in enumerate(DISPLAY_HEADERS, start=1):
        cell = sheet.cell(row=2, column=col_idx, value=header)
        cell.font = styles["header_font"]
        cell.fill = hdr_fill
        cell.alignment = styles["center"]
        cell.border = styles["thin_border"]
    sheet.row_dimensions[2].height = 18

    alt = styles["net_alt_fill"] if is_net else styles["alt_fill"]
    for row_idx, angle in enumerate(ALL_ANGLES, start=3):
        vals = angle_data.get(angle, {})
        fill = alt if (row_idx % 2 == 0) else None

        ac = sheet.cell(row=row_idx, column=1, value=angle)
        ac.font = styles["loc_font"]
        ac.alignment = styles["center"]
        ac.border = styles["thin_border"]
        if fill:
            ac.fill = fill

        for col_idx, col_key in enumerate(COLUMNS, start=2):
            val = vals.get(col_key)
            cell = sheet.cell(row=row_idx, column=col_idx, value=val)
            cell.font = styles["data_font"]
            cell.alignment = styles["center"]
            cell.border = styles["thin_border"]
            if fill:
                cell.fill = fill
            cell.number_format = "0.000" if col_key in ("Vh", "TI") else "0.0"

    for i, w in enumerate([12, 10, 10, 10, 10], start=1):
        sheet.column_dimensions[get_column_letter(i)].width = w
    sheet.freeze_panes = "A3"


def write_stab_excel(data, output_path, target_h, stab_name):
    """One Excel file per stab class (unchanged behaviour)."""
    wb = Workbook()
    wb.remove(wb.active)
    styles = make_styles()

    for loc_name in sorted(data.keys()):
        write_data_sheet(
            wb,
            sheet_title=f"{stab_name} - {loc_name}",
            title_text=f"{loc_name}  |  {stab_name}  |  H = {target_h} m",
            angle_data=data[loc_name],
            styles=styles,
        )

    if not wb.sheetnames:
        print(f"    No data to write for {stab_name}.")
        return False

    wb.save(output_path)
    print(f"    Saved: {os.path.basename(output_path)}  ({len(wb.sheetnames)} sheet(s))")
    return True


def write_combined_excel(all_stab_data, net_data, weights, output_path, target_h):
    """One combined Excel: all stab sheets in order, then Net sheets at the end."""
    wb = Workbook()
    wb.remove(wb.active)
    styles = make_styles()

    total_weight = sum(weights.values())
    weight_note = "  |  Weights: " + ", ".join(
        f"{s}={w/total_weight*100:.1f}%" for s, w in sorted(weights.items())
    )

    # All stab sheets
    for stab_name in sorted(all_stab_data.keys(), key=lambda x: int(re.search(r'\d+', x).group())):
        data = all_stab_data[stab_name]
        for loc_name in sorted(data.keys()):
            write_data_sheet(
                wb,
                sheet_title=f"{stab_name} - {loc_name}",
                title_text=f"{loc_name}  |  {stab_name}  |  H = {target_h} m",
                angle_data=data[loc_name],
                styles=styles,
            )

    # Net sheets (one per location), visually distinct
    for loc_name in sorted(net_data.keys()):
        write_data_sheet(
            wb,
            sheet_title=f"Net - {loc_name}",
            title_text=f"{loc_name}  |  Net (weighted avg)  |  H = {target_h} m{weight_note}",
            angle_data=net_data[loc_name],
            styles=styles,
            is_net=True,
        )

    wb.save(output_path)
    stab_sheets = [s for s in wb.sheetnames if not s.startswith("Net")]
    net_sheets  = [s for s in wb.sheetnames if s.startswith("Net")]
    print(f"    Saved: {os.path.basename(output_path)}")
    print(f"      {len(stab_sheets)} stab sheet(s), {len(net_sheets)} Net sheet(s): {', '.join(net_sheets)}")


# ── Folder discovery ──────────────────────────────────────────────────────────

def find_stab_folders(parent_folder):
    stab_folders = []
    for entry in os.scandir(parent_folder):
        if entry.is_dir() and re.match(r'(?i)stab_\d+', entry.name):
            stab_folders.append((entry.name, entry.path))
    stab_folders.sort(key=lambda x: int(re.search(r'\d+', x[0]).group()))
    return stab_folders


# ── Main ──────────────────────────────────────────────────────────────────────

def main():
    print("=== Wind Data Extractor ===\n")

    # Parent folder
    while True:
        parent = input("Parent folder path (containing Stab_X subfolders): ").strip().strip('"').strip("'")
        if os.path.isdir(parent):
            break
        print(f"  Folder not found: '{parent}' — please try again.\n")

    stab_folders = find_stab_folders(parent)
    if not stab_folders:
        print(f"\n  No Stab_X folders found in: {parent}")
        input("\nPress Enter to exit.")
        return

    stab_names = [n for n, _ in stab_folders]
    print(f"\n  Found {len(stab_folders)} stability class(es): {', '.join(stab_names)}")

    # Height
    while True:
        h_raw = input("\nHeight H (m) to extract (e.g. 120): ").strip()
        try:
            target_h = float(h_raw)
            break
        except ValueError:
            print("  Please enter a valid number.\n")

    # Weights — enter one per stab, don't need to sum to 1
    print(f"\nEnter a weight for each stability class (any scale, e.g. 1 1 2 or 25 25 50).")
    weights = {}
    for stab_name in stab_names:
        while True:
            w_raw = input(f"  Weight for {stab_name}: ").strip()
            try:
                weights[stab_name] = float(w_raw)
                break
            except ValueError:
                print("  Please enter a valid number.")

    total_w = sum(weights.values())
    print(f"\n  Normalised weights:")
    for s, w in weights.items():
        print(f"    {s}: {w/total_w*100:.1f}%")

    # Output folder
    default_out = parent
    out_raw = input(f"\nOutput folder (press Enter for: {default_out}): ").strip().strip('"').strip("'")
    out_folder = out_raw if out_raw else default_out
    os.makedirs(out_folder, exist_ok=True)

    # ── Process ──
    print(f"\nProcessing H = {target_h} m ...\n")

    all_stab_data = {}
    for stab_name, stab_path in stab_folders:
        print(f"  [{stab_name}]  {stab_path}")
        data = collect_data(stab_path, target_h)
        if not data:
            print(f"    No data found — skipping.\n")
            continue
        all_stab_data[stab_name] = data

        # Individual stab Excel
        out_path = os.path.join(out_folder, f"{stab_name}_H{int(target_h)}m.xlsx")
        write_stab_excel(data, out_path, target_h, stab_name)
        print()

    if not all_stab_data:
        print("No data extracted at all. Check folder structure and H value.")
        input("\nPress Enter to exit.")
        return

    # Combined Excel with Net sheet
    print("  [Combined + Net]")
    net_data = compute_net_data(all_stab_data, weights)
    combined_path = os.path.join(out_folder, f"Combined_Net_H{int(target_h)}m.xlsx")
    write_combined_excel(all_stab_data, net_data, weights, combined_path, target_h)

    print(f"\nDone! All files saved to: {out_folder}")
    input("\nPress Enter to exit.")


if __name__ == "__main__":
    main()